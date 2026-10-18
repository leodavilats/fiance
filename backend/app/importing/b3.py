from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.pregao import deslocar_pregoes
from app.ledger import LedgerEntry, LedgerError, TransactionKind

from .parser import MAX_ROWS, ImportIssue, parse_day, parse_decimal, parse_symbol

FORMATO = "extrato de movimentação da B3"

MAX_BYTES = 2_000_000

CICLO_DE_LIQUIDACAO = 2

_COLUNAS = {
    "entrada/saida": "direction",
    "data": "day",
    "movimentacao": "movement",
    "produto": "product",
    "quantidade": "quantity",
    "preco unitario": "price",
    "valor da operacao": "value",
}

_LIQUIDACAO = "transferencia - liquidacao"
_ATUALIZACAO = "atualizacao"
_PROVENTOS = {
    "dividendo": "dividendo",
    "juros sobre capital proprio": "jcp",
    "rendimento": "rendimento",
}
_EVENTOS = {
    "desdobro",
    "desdobramento",
    "grupamento",
    "bonificacao em ativos",
    "incorporacao",
    "fracao em ativos",
    "leilao de fracao",
}

_FRACIONARIO = re.compile(r"^([A-Z]{4}[0-9]{1,2})F$")


@dataclass
class B3Operation:
    line: int
    entry: LedgerEntry
    settled_on: str


@dataclass
class B3Dividend:
    line: int
    ticker: str
    paid_at: str
    amount: float
    kind: str
    quantity: float
    price: float


@dataclass(frozen=True)
class B3Custody:
    line: int
    symbol: str
    on: str
    quantity: float


@dataclass
class ParsedB3:
    operations: list[B3Operation] = field(default_factory=list)
    dividends: list[B3Dividend] = field(default_factory=list)
    custody: list[B3Custody] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)


def _chave(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def _dia(valor: object) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return parse_day(str(valor or ""), 0)


def _br(dia: str) -> str:
    return date.fromisoformat(dia).strftime("%d/%m/%Y")


def _numero(valor: object, nome: str) -> float:
    if isinstance(valor, bool):
        raise ValueError(f"{nome} não é um número")
    if isinstance(valor, int | float):
        return float(valor)
    texto = str(valor or "").strip()
    if not texto or texto == "-":
        raise ValueError(f"{nome} vazio")
    return parse_decimal(texto, 0, nome)


def _ticker(produto: str) -> str | None:
    codigo = produto.split(" - ", 1)[0].strip().upper()
    fracionario = _FRACIONARIO.fullmatch(codigo)
    if fracionario:
        codigo = fracionario.group(1)
    try:
        return parse_symbol(codigo)
    except ValueError:
        return None


def _ler_planilha(conteudo: bytes, result: ParsedB3) -> list[tuple] | None:
    if len(conteudo) > MAX_BYTES:
        result.issues.append(
            ImportIssue(line=0, message="O arquivo passa de 2 MB. Baixe um período menor.")
        )
        return None
    try:
        livro = load_workbook(io.BytesIO(conteudo), data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        result.issues.append(
            ImportIssue(
                line=0,
                message=(
                    "Não consegui abrir o arquivo como Excel. Baixe o extrato de movimentação "
                    "em Excel (.xlsx) na Área do Investidor da B3."
                ),
            )
        )
        return None
    try:
        return list(livro.worksheets[0].iter_rows(values_only=True))
    finally:
        livro.close()


def parse_b3_movimentacao(conteudo: bytes) -> ParsedB3:
    result = ParsedB3()

    linhas = _ler_planilha(conteudo, result)
    if linhas is None:
        return result
    if not linhas:
        result.issues.append(ImportIssue(line=0, message="A planilha está vazia."))
        return result

    cabecalho = [_COLUNAS.get(_chave(celula)) for celula in linhas[0]]
    if any(nome not in cabecalho for nome in _COLUNAS.values()):
        result.issues.append(
            ImportIssue(
                line=1,
                message=(
                    "Não parece o extrato de movimentação da B3: o cabeçalho precisa ter "
                    "Entrada/Saída, Data, Movimentação, Produto, Quantidade, Preço unitário e "
                    "Valor da Operação."
                ),
                raw=" | ".join(str(c or "") for c in linhas[0]),
            )
        )
        return result

    corpo = [
        (numero_da_linha, linha)
        for numero_da_linha, linha in enumerate(linhas[1:], start=2)
        if any(celula not in (None, "") for celula in linha)
    ]
    if len(corpo) > MAX_ROWS:
        result.issues.append(
            ImportIssue(
                line=0,
                message=(
                    f"São {len(corpo)} linhas e o limite é {MAX_ROWS}. Baixe o extrato em "
                    "períodos menores."
                ),
            )
        )
        return result

    for numero_da_linha, linha in corpo:
        valores = {
            nome: linha[posicao]
            for posicao, nome in enumerate(cabecalho)
            if nome and posicao < len(linha)
        }
        bruta = " | ".join("" if c is None else str(c) for c in linha)
        try:
            _consumir(valores, numero_da_linha, bruta, result)
        except (ValueError, LedgerError) as exc:
            result.issues.append(ImportIssue(line=numero_da_linha, message=str(exc), raw=bruta))

    return result


def _consumir(valores: dict, linha: int, bruta: str, result: ParsedB3) -> None:
    movimento = _chave(valores.get("movement"))
    produto = str(valores.get("product") or "").strip()
    rotulo = str(valores.get("movement") or "").strip()

    if movimento == _ATUALIZACAO:
        ticker = _ticker(produto)
        if ticker is not None:
            result.custody.append(
                B3Custody(
                    line=linha,
                    symbol=ticker,
                    on=_dia(valores.get("day")),
                    quantity=_numero(valores.get("quantity"), "quantidade"),
                )
            )
        return

    if movimento != _LIQUIDACAO and movimento not in _PROVENTOS:
        dica = (
            "lance o evento à mão no razão, com a proporção."
            if movimento in _EVENTOS
            else "se ela mudou a sua posição, lance à mão no razão."
        )
        result.warnings.append(
            ImportIssue(line=linha, message=f"“{rotulo}” fica de fora: {dica}", raw=bruta)
        )
        return

    ticker = _ticker(produto)
    if ticker is None:
        result.warnings.append(
            ImportIssue(
                line=linha,
                message=(
                    f"“{produto}” não tem código de negociação da B3 e fica de fora. "
                    "Renda fixa se cadastra em Renda fixa."
                ),
                raw=bruta,
            )
        )
        return

    direcao = _chave(valores.get("direction"))
    if direcao not in {"credito", "debito"}:
        raise ValueError(f"Entrada/Saída {valores.get('direction')!r} não é Crédito nem Débito.")
    dia = _dia(valores.get("day"))

    if movimento in _PROVENTOS:
        if direcao != "credito":
            result.warnings.append(
                ImportIssue(
                    line=linha,
                    message=f"{rotulo} de {ticker} veio como débito e fica de fora.",
                    raw=bruta,
                )
            )
            return
        valor = _numero(valores.get("value"), "valor da operação")
        if valor <= 0:
            raise ValueError("Provento sem valor positivo.")
        result.dividends.append(
            B3Dividend(
                line=linha,
                ticker=ticker,
                paid_at=dia,
                amount=round(valor, 2),
                kind=_PROVENTOS[movimento],
                quantity=_numero(valores.get("quantity"), "quantidade"),
                price=_numero(valores.get("price"), "preço unitário"),
            )
        )
        return

    pregao = deslocar_pregoes(date.fromisoformat(dia), -CICLO_DE_LIQUIDACAO).isoformat()
    result.operations.append(
        B3Operation(
            line=linha,
            settled_on=dia,
            entry=LedgerEntry(
                kind=TransactionKind.BUY if direcao == "credito" else TransactionKind.SELL,
                symbol=ticker,
                traded_on=pregao,
                quantity=_numero(valores.get("quantity"), "quantidade"),
                price=round(_numero(valores.get("price"), "preço unitário"), 8),
                note=f"Extrato da B3: liquidada em {_br(dia)}.",
            ),
        )
    )
